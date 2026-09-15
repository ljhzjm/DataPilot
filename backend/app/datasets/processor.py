import asyncio
import csv
import shutil
from collections.abc import Sequence
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Protocol
from uuid import UUID

import duckdb
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from app.tools.duckdb_engine import ColumnInfo, DuckDBAnalyticsEngine

SUPPORTED_EXTENSIONS = {".csv", ".tsv", ".xlsx"}


class AsyncUpload(Protocol):
    async def read(self, size: int = -1) -> bytes: ...


class ProcessedDataset(BaseModel):
    id: UUID
    name: str
    original_filename: str
    table_name: str
    file_type: str
    file_size: int = Field(ge=1)
    original_path: Path
    parquet_path: Path
    row_count: int = Field(ge=0)
    columns: list[ColumnInfo]


class DatasetProcessingError(ValueError):
    pass


class DatasetProcessor:
    def __init__(
        self,
        *,
        engine: DuckDBAnalyticsEngine,
        upload_root: Path,
        max_upload_bytes: int = 20 * 1024 * 1024,
        max_rows: int = 500_000,
        max_columns: int = 500,
    ) -> None:
        self._engine = engine
        self._upload_root = upload_root.resolve()
        self._max_upload_bytes = max_upload_bytes
        self._max_rows = max_rows
        self._max_columns = max_columns

    async def process(
        self,
        *,
        dataset_id: UUID,
        filename: str,
        upload: AsyncUpload,
    ) -> ProcessedDataset:
        safe_filename = Path(filename).name
        extension = Path(safe_filename).suffix.casefold()
        if extension not in SUPPORTED_EXTENSIONS:
            raise DatasetProcessingError("Unsupported file type. Allowed: .csv, .tsv, .xlsx.")

        table_name = f"dataset_{dataset_id.hex[:12]}"
        dataset_dir = self._upload_root / str(dataset_id)
        dataset_dir.mkdir(parents=True, exist_ok=False)
        original_path = dataset_dir / f"source{extension}"
        parquet_path = dataset_dir / "data.parquet"

        try:
            file_size = await self._write_upload(upload, original_path)
            csv_source = original_path
            if extension == ".xlsx":
                csv_source = dataset_dir / "normalized.csv"
                await asyncio.to_thread(
                    self._xlsx_to_csv,
                    original_path,
                    csv_source,
                )

            await asyncio.to_thread(
                self._convert_to_parquet,
                csv_source,
                parquet_path,
            )
            row_count = await asyncio.to_thread(
                self._count_parquet_rows,
                parquet_path,
            )
            await asyncio.to_thread(
                self._engine.register_parquet,
                table_name,
                parquet_path,
            )
            schema = await asyncio.to_thread(
                self._engine.get_schema,
                table_name,
            )
            if len(schema.columns) > self._max_columns:
                raise DatasetProcessingError(f"Dataset exceeds {self._max_columns} columns.")
            return ProcessedDataset(
                id=dataset_id,
                name=Path(safe_filename).stem[:160] or "dataset",
                original_filename=safe_filename,
                table_name=table_name,
                file_type=extension.lstrip("."),
                file_size=file_size,
                original_path=original_path,
                parquet_path=parquet_path,
                row_count=row_count,
                columns=schema.columns,
            )
        except Exception:
            self.remove(dataset_id, table_name)
            raise

    def remove(self, dataset_id: UUID, table_name: str) -> None:
        self._engine.drop_table(table_name)
        shutil.rmtree(self._upload_root / str(dataset_id), ignore_errors=True)

    async def _write_upload(self, upload: AsyncUpload, target: Path) -> int:
        size = 0
        with target.open("wb") as output:
            while True:
                chunk = await upload.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > self._max_upload_bytes:
                    raise DatasetProcessingError(f"File exceeds {self._max_upload_bytes} bytes.")
                await asyncio.to_thread(output.write, chunk)
        if size == 0:
            raise DatasetProcessingError("Uploaded file is empty.")
        return size

    def _xlsx_to_csv(self, source: Path, target: Path) -> None:
        workbook = load_workbook(
            filename=source,
            read_only=True,
            data_only=True,
        )
        try:
            worksheet = workbook.active
            if worksheet is None:
                raise DatasetProcessingError("Excel workbook has no active worksheet.")
            rows = worksheet.iter_rows(values_only=True)
            try:
                raw_headers = next(rows)
            except StopIteration as exc:
                raise DatasetProcessingError("Excel worksheet is empty.") from exc

            headers = _unique_headers(raw_headers)
            if len(headers) > self._max_columns:
                raise DatasetProcessingError(f"Dataset exceeds {self._max_columns} columns.")

            row_count = 0
            with target.open("w", encoding="utf-8", newline="") as output:
                writer = csv.writer(output)
                writer.writerow(headers)
                for row in rows:
                    if row_count >= self._max_rows:
                        raise DatasetProcessingError(f"Dataset exceeds {self._max_rows} rows.")
                    values = [_normalize_cell(value) for value in row[: len(headers)]]
                    if len(values) < len(headers):
                        values.extend([""] * (len(headers) - len(values)))
                    writer.writerow(values)
                    row_count += 1
            if row_count == 0:
                raise DatasetProcessingError("Excel worksheet has no data rows.")
        finally:
            workbook.close()

    def _convert_to_parquet(
        self,
        source: Path,
        target: Path,
    ) -> None:
        connection = duckdb.connect()
        try:
            relation = connection.read_csv(
                str(source),
                header=True,
                sample_size=-1,
            )
            relation.write_parquet(str(target), compression="zstd")
        finally:
            connection.close()

    def _count_parquet_rows(self, parquet_path: Path) -> int:
        connection = duckdb.connect()
        try:
            row = connection.execute(
                "SELECT COUNT(*) FROM read_parquet(?)",
                [str(parquet_path)],
            ).fetchone()
        finally:
            connection.close()
        if row is None:
            return 0
        row_count = int(row[0])
        if row_count > self._max_rows:
            raise DatasetProcessingError(f"Dataset exceeds {self._max_rows} rows.")
        return row_count


def _unique_headers(values: Sequence[Any]) -> list[str]:
    headers: list[str] = []
    used: set[str] = set()
    for index, value in enumerate(values):
        base = str(value).strip() if value is not None else ""
        base = base or f"column_{index + 1}"
        candidate = base
        suffix = 2
        while candidate.casefold() in used:
            candidate = f"{base}_{suffix}"
            suffix += 1
        used.add(candidate.casefold())
        headers.append(candidate)
    return headers


def _normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)
